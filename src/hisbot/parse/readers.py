"""文件读取器（文档 4.7 / 表 15）。

xlsx(openpyxl) / xls(xlrd)：合并单元格还原、多级表头展开、日期序列号转换；
csv：UTF-8 / UTF-8-BOM / GBK 自动判定 + 分隔符嗅探；
pdf：优先表格线抽取，无线则按字符坐标聚类还原行列；无文本层抛 PRS-002。
统一产出 RawTable（单层规范表头 + 字符串二维数据）。
"""
from __future__ import annotations

import logging

_log = logging.getLogger("hisbot")

import datetime as _dt
from dataclasses import dataclass, field
from typing import Any


@dataclass
class RawTable:
    headers: list[str]
    rows: list[list[str]]
    header_row_count: int = 1
    parser: str = ""
    confidence: float = 1.0
    extras: dict = field(default_factory=dict)


# --------------------------------------------------------------------------- #
# 单元格标准化
# --------------------------------------------------------------------------- #
def cell_to_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S") if v.time() != _dt.time(0) \
            else v.strftime("%Y-%m-%d")
    if isinstance(v, _dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(v)
    return str(v).strip()


def _matrix_openpyxl(ws) -> list[list[Any]]:
    data = [[c.value for c in row] for row in ws.iter_rows()]
    max_c = max([ws.max_column or 0] + [len(r) for r in data])
    for r in data:
        r += [None] * (max_c - len(r))
    # 合并单元格：用左上角值填充整个合并区
    for rng in ws.merged_cells.ranges:
        top = ws.cell(rng.min_row, rng.min_col).value
        for row_idx in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if row_idx - 1 < len(data) and c - 1 < max_c:
                    data[row_idx - 1][c - 1] = top
    return data


def _matrix_xlrd(book, sheet) -> list[list[Any]]:
    def val(r, c):
        cell = sheet.cell(r, c)
        if cell.ctype == 3:  # XL_CELL_DATE
            try:
                from xlrd import xldate_as_datetime
                return xldate_as_datetime(cell.value, book.datemode)
            except Exception:
                return cell.value
        return cell.value
    data = [[val(r, c) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)]
    # xlrd 合并区 (rlo,rhi,clo,chi)
    for (rlo, rhi, clo, chi) in sheet.merged_cells:
        top = data[rlo][clo]
        for r in range(rlo, rhi):
            for c in range(clo, chi):
                data[r][c] = top
    return data


def _trim_matrix(mat: list[list[str]]) -> list[list[str]]:
    """去除全空行、全空列。"""
    mat = [r for r in mat if any(str(c).strip() for c in r)]
    if not mat:
        return mat
    ncol = max(len(r) for r in mat)
    for r in mat:
        r += [""] * (ncol - len(r))
    keep = [j for j in range(ncol)
            if any(str(mat[i][j]).strip() for i in range(len(mat)))]
    return [[r[j] for j in keep] for r in mat]


def _is_data_like(s: str) -> bool:
    s = str(s).strip().replace(",", "").replace("，", "")
    if not s:
        return False
    try:
        float(s); return True
    except Exception:
        _log.warning("忽略异常 @%s", __name__, exc_info=True)
    import re
    if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", s):
        return True
    return False


def build_headers(mat: list[list[str]],
                  forced_header_rows: int | None = None
                  ) -> tuple[list[str], int]:
    """结构探测 + 多级表头展开。

    默认单级表头（第 0 行）；仅当外部依据表头区横向合并单元格明确给出
    forced_header_rows>1 时按多级处理（避免把含数字的数据行误判成表头）。
    多级表头纵向拼接为"一级_二级"，重复列名自动加序号。
    """
    if not mat:
        return [], 0
    ncol = max(len(r) for r in mat)
    for r in mat:
        r += [""] * (ncol - len(r))
    hdr_rows = max(1, int(forced_header_rows or 1))
    if forced_header_rows is None:
        hdr_rows = 1
    hdr_rows = min(hdr_rows, len(mat))

    grid = [[str(mat[r][c]).strip() if c < len(mat[r]) else ""
             for c in range(ncol)] for r in range(hdr_rows)]
    headers = []
    for j in range(ncol):
        parts: list[str] = []
        for i in range(hdr_rows):
            v = grid[i][j]
            if not v and i > 0:
                v = grid[0][j]  # 向下继承一级
            if v and (not parts or v != parts[-1]):
                parts.append(v)
        headers.append("_".join(parts))
    # 去重列名
    seen: dict[str, int] = {}
    out: list[str] = []
    for h in headers:
        h = h or f"未命名列{len(out) + 1}"
        if h in seen:
            seen[h] += 1
            out.append(f"{h}_{seen[h]}")
        else:
            seen[h] = 1
            out.append(h)
    return out, hdr_rows


# --------------------------------------------------------------------------- #
# xlsx / xls
# --------------------------------------------------------------------------- #
def read_xlsx(path: str) -> RawTable:
    from openpyxl import load_workbook
    # 需要 merged_cells.ranges 还原合并单元格，故不使用 read_only 模式
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb.active
    hrows = 1
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 1 and rng.max_col > rng.min_col:
            # 第 1 行存在横向合并分组标题 => 字段名位于第 2 行，至少两级
            hrows = max(hrows, 2, rng.max_row)
    raw = _matrix_openpyxl(ws)
    wb.close()
    return _matrix_to_table(raw, "xlsx", hrows)


def read_xls(path: str) -> RawTable:
    import xlrd
    book = xlrd.open_workbook(path)
    sheet = book.sheet_by_index(0)
    hrows = 1
    for (rlo, rhi, clo, chi) in sheet.merged_cells:
        if rlo == 0 and chi - clo > 1:
            hrows = max(hrows, 2, rhi)
    raw = _matrix_xlrd(book, sheet)
    return _matrix_to_table(raw, "xls", hrows)


def _matrix_to_table(raw: list[list[Any]], parser: str,
                     header_rows: int | None = None) -> RawTable:
    mat = [[cell_to_str(v) for v in row] for row in raw]
    mat = _trim_matrix(mat)
    if not mat:
        return RawTable([], [], 0, parser)
    headers, hcnt = build_headers(mat, header_rows)
    body = mat[hcnt:]
    return RawTable(headers, body, hcnt, parser)


# --------------------------------------------------------------------------- #
# csv
# --------------------------------------------------------------------------- #
def detect_encoding(path: str) -> str:
    with open(path, "rb") as f:
        raw = f.read(65536)
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if raw.startswith(b"\xff\xfe"):
        return "utf-16-le"
    if raw.startswith(b"\xfe\xff"):
        return "utf-16-be"
    best, best_score = "utf-8", -1
    for enc in ("utf-8", "gb18030", "gbk", "big5", "latin-1"):
        try:
            txt = raw.decode(enc, errors="strict")
            # 越少替换/常见中文乱码字符越好
            bad = txt.count("�") + txt.count("\ufffd")
            cjk = sum(1 for ch in txt if "\u4e00" <= ch <= "\u9fff")
            score = cjk * 2 - bad * 100
            if score > best_score:
                best, best_score = enc, score
        except Exception:
            continue
    return best


def read_csv(path: str, encoding: str = "auto",
             delimiter: str | None = None) -> RawTable:
    import csv as _csv
    enc = detect_encoding(path) if encoding in ("auto", "", None) \
        else encoding
    with open(path, encoding=enc, errors="replace", newline="") as f:
        text = f.read()
    if delimiter is None:
        try:
            dialect = _csv.Sniffer().sniff(text[:65536],
                                           delimiters=",;\t|，；")
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","
    import io
    rows = [[(c or "").strip() for c in r]
            for r in _csv.reader(io.StringIO(text), delimiter=delimiter)]
    rows = _trim_matrix(rows)
    if not rows:
        return RawTable([], [], 0, "csv")
    headers, hcnt = build_headers(rows)
    return RawTable(headers, rows[hcnt:], hcnt, "csv",
                    extras={"encoding": enc, "delimiter": delimiter})


# --------------------------------------------------------------------------- #
# pdf
# --------------------------------------------------------------------------- #
def read_pdf(path: str) -> RawTable:
    import pdfplumber
    matrix: list[list[str]] = []
    has_text = False
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                has_text = True
            tables = page.extract_tables()
            got = False
            for tb in tables:
                if tb:
                    for row in tb:
                        matrix.append([(c or "").replace("\n", " ").strip()
                                       for c in row])
                    got = True
            if not got:
                matrix.extend(_cluster_words(page))
    if not has_text and not matrix:
        from ..core.errors import PdfNoTextLayerError
        raise PdfNoTextLayerError("PDF 无文本层（扫描件），不支持解析")
    matrix = _trim_matrix(matrix)
    if not matrix:
        return RawTable([], [], 0, "pdf", 0.5)
    headers, hcnt = build_headers(matrix)
    return RawTable(headers, matrix[hcnt:], hcnt, "pdf", 0.9)


def _cluster_words(page, row_tol: float = 3.0) -> list[list[str]]:
    """无表格线时按字符坐标聚类还原行列。"""
    words = page.extract_words(use_text_flow=False, keep_blank_chars=False)
    if not words:
        return []
    # 行聚类：top 排序后按容差归并
    words = sorted(words, key=lambda w: (w["top"], w["x0"]))
    rows: list[list[dict]] = []
    for w in words:
        placed = False
        for row in rows:
            if abs(w["top"] - row[0]["top"]) <= max(
                    row_tol, w["height"] * 0.6):
                row.append(w); placed = True; break
        if not placed:
            rows.append([w])
    # 列聚类：收集所有词的 x 中心
    xs = sorted((w["x0"] + w["x1"]) / 2 for row in rows for w in row)
    centers: list[float] = []
    for x in xs:
        if not centers or abs(x - centers[-1]) > 12:
            centers.append(x)
        else:
            centers[-1] = (centers[-1] + x) / 2

    out: list[list[str]] = []
    for row in rows:
        cells = [""] * len(centers)
        for w in row:
            cx = (w["x0"] + w["x1"]) / 2
            ci = min(range(len(centers)),
                     key=lambda i: abs(cx - centers[i]))
            cells[ci] = (cells[ci] + w["text"]).strip()
        out.append(cells)
    return out


def read_file(path: str, ext: str | None = None,
              encoding: str = "auto") -> RawTable:
    import os
    ext = (ext or os.path.splitext(path)[1]).lower()
    if ext == ".xlsx":
        return read_xlsx(path)
    if ext == ".xls":
        return read_xls(path)
    if ext == ".csv":
        return read_csv(path, encoding)
    if ext == ".pdf":
        return read_pdf(path)
    from ..core.errors import UnsupportedFileTypeError
    raise UnsupportedFileTypeError(f"不支持的文件类型：{ext}")
